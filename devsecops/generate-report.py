import argparse
import json
import os
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPORT_DIR = Path(__file__).resolve().parent / "reports"
DEFAULT_INPUT_FILE = REPORT_DIR / "vulnerability-report.sarif.json"
DEFAULT_OUTPUT_FILE = REPORT_DIR / "vulnerability-report.xlsx"
DEFAULT_RECOMMENDATIONS_FILE = REPORT_DIR / "base-image-recommendations.txt"
REPORT_COLUMNS = [
    "CVE ID",
    "Severity",
    "Package",
    "Installed Version",
    "Affected Range",
    "Fixed Version",
    "Fix Available",
    "CVSS Score",
    "EPSS Score",
    "EPSS Percentile",
    "Advisory URL",
    "Location",
]


def load_report(input_file):
    """Load a vulnerability report from JSON."""
    if not input_file.exists():
        raise FileNotFoundError(f"Input file '{input_file}' not found.")

    with input_file.open("r", encoding="utf-8") as file:
        return json.load(file)


def _split_recommendation_line(line):
    return [part.strip() for part in line.split(chr(0x2502))]


def _recommendation_vulnerabilities(line):
    match = re.search(
        r"(\d+)C\s+(\d+)H\s+(\d+)M\s+(\d+)L(?:\s+(\d+)\?)?",
        line,
    )
    if not match:
        return None
    return [int(value or 0) for value in match.groups()]


def load_recommendations(recommendations_file):
    """Parse Docker Scout recommendations into structured worksheet data."""
    if not recommendations_file.exists():
        return {
            "overview": {
                "Status": "No Docker Scout recommendations file was generated."
            },
            "candidates": [],
        }

    overview = {}
    candidates = []
    section = ""
    known_fields = {
        "Target": "Target",
        "digest": "Digest",
        "Name": "Recommended Base Image",
        "Vulnerabilities": "Current Vulnerabilities",
        "Pushed": "Pushed",
        "Size": "Size",
        "Packages": "Packages",
        "Flavor": "Flavor",
        "Runtime": "Runtime",
    }

    current_candidate = None
    for raw_line in recommendations_file.read_text(
        encoding="utf-8", errors="replace"
    ).splitlines():
        stripped = raw_line.strip()
        if not stripped:
            continue

        if stripped in {"Refresh base image", "Change base image"}:
            section = stripped
            current_candidate = None
            continue

        if stripped.lower().startswith("base image is"):
            overview["Current Base Image"] = stripped.split("is", 1)[1].strip()
            continue

        parts = _split_recommendation_line(raw_line)
        first_cell = parts[0].strip() if parts else ""

        if first_cell in known_fields and len(parts) > 1:
            overview[known_fields[first_cell]] = parts[1]

        vulnerability_counts = _recommendation_vulnerabilities(raw_line)
        is_tag = bool(
            re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", first_cell)
        )
        if not vulnerability_counts or not section or not is_tag:
            if current_candidate and chr(0x2502) in raw_line:
                if chr(0x2500) in raw_line:
                    continue
                first_detail = first_cell
                second_detail = parts[1] if len(parts) > 1 else ""
                detail_parts = []
                if first_detail and first_detail not in {
                    "Benefits:",
                    "Image details:",
                }:
                    detail_parts.append(first_detail)
                if second_detail and second_detail not in {
                    "Benefits:",
                    "Image details:",
                }:
                    detail_parts.append(second_detail)
                if detail_parts:
                    current_candidate["Details"] += "\n" + ": ".join(
                        detail_parts
                    )
            continue

        pushed = parts[2] if len(parts) > 2 else ""
        detail = parts[1] if len(parts) > 1 else ""
        current_candidate = {
                "Recommendation": section,
                "Tag": first_cell,
                "Details": detail,
                "Pushed": pushed,
                "Critical": vulnerability_counts[0],
                "High": vulnerability_counts[1],
                "Medium": vulnerability_counts[2],
                "Low": vulnerability_counts[3],
                "Unspecified": vulnerability_counts[4],
            }
        candidates.append(current_candidate)

    if not overview:
        overview["Status"] = "Docker Scout returned no structured recommendation metadata."

    return {"overview": overview, "candidates": candidates}


def _text(value):
    if isinstance(value, dict):
        return value.get("text", value.get("name", ""))
    return value if isinstance(value, str) else ""


def _number(value):
    if value in (None, "", "not available", "unknown"):
        return ""
    try:
        return float(value)
    except (TypeError, ValueError):
        return value


def _parse_message(message):
    """Parse Docker Scout's aligned key/value message text."""
    fields = {}
    for line in message.splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        normalized_key = re.sub(
            r"[^a-z0-9]+", "_", key.strip().lower()
        ).strip("_")
        fields[normalized_key] = value.strip()
    return fields


def _purl_details(purls):
    """Extract package name, installed version, and the original PURL."""
    if not isinstance(purls, list) or not purls:
        return "", "", ""

    purl = str(purls[0])
    without_query = purl.split("?", 1)[0]
    if "@" not in without_query:
        return without_query.rsplit("/", 1)[-1], "", purl

    package_path, version = without_query.rsplit("@", 1)
    return package_path.rsplit("/", 1)[-1], version, purl


def _locations(result):
    locations = []
    for location in result.get("locations", []):
        uri = (
            location.get("physicalLocation", {})
            .get("artifactLocation", {})
            .get("uri", "")
        )
        if uri and uri not in locations:
            locations.append(uri)
    return "; ".join(locations)


def extract_sarif_vulnerabilities(data):
    """Convert Docker Scout SARIF results into report rows."""
    vulnerabilities = []

    for run in data.get("runs", []):
        rules = {
            rule.get("id", ""): rule
            for rule in run.get("tool", {}).get("driver", {}).get("rules", [])
        }

        for result in run.get("results", []):
            rule_id = result.get("ruleId", "")
            rule = rules.get(rule_id, {})
            properties = {}
            properties.update(rule.get("properties", {}))
            properties.update(result.get("properties", {}))

            message = _text(result.get("message", ""))
            message_fields = _parse_message(message)
            package_from_purl, version_from_purl, _ = _purl_details(
                properties.get("purls", [])
            )
            package = (
                properties.get("package")
                or package_from_purl
                or message_fields.get("package", "")
            )
            if str(package).startswith("pkg:"):
                package, _, _ = _purl_details([package])

            severity = (
                properties.get("cvssV3_severity")
                or properties.get("severity")
                or (properties.get("tags") or [""])[0]
                or message_fields.get("severity", "UNKNOWN")
            )

            fixed_version = properties.get(
                "fixed_version", message_fields.get("fixed_version", "")
            )
            fixed_version = str(fixed_version)
            fix_available = fixed_version.lower() not in {
                "",
                "not fixed",
                "not available",
                "unknown",
            }

            vulnerabilities.append(
                {
                    "cve_id": rule_id,
                    "severity": str(severity).upper(),
                    "package": package,
                    "installed_version": properties.get(
                        "installed_version",
                        properties.get(
                            "version",
                            message_fields.get(
                                "installed_version", version_from_purl
                            ),
                        ),
                    ),
                    "affected_range": properties.get(
                        "affected_version",
                        properties.get(
                            "affected_range",
                            message_fields.get("affected_range", ""),
                        ),
                    ),
                    "fixed_version": fixed_version,
                    "fix_available": fix_available,
                    "cvss_score": _number(
                        properties.get("security-severity", "")
                    ),
                    "epss_score": _number(message_fields.get("epss_score", "")),
                    "epss_percentile": _number(
                        message_fields.get("epss_percentile", "")
                    ),
                    "advisory_url": rule.get("helpUri", ""),
                    "location": _locations(result),
                }
            )

    return vulnerabilities


def extract_vulnerabilities(data):
    """Extract vulnerability records from SARIF or compatible JSON."""
    if isinstance(data, dict) and "runs" in data:
        return extract_sarif_vulnerabilities(data)

    vulnerabilities = []
    if isinstance(data, list):
        vulnerabilities = data
    elif isinstance(data, dict):
        if "vulnerabilities" in data:
            vulnerabilities = data["vulnerabilities"]
        elif "matches" in data:
            vulnerabilities = data["matches"]
        elif "packages" in data:
            for package in data["packages"]:
                for vulnerability in package.get("vulnerabilities", []):
                    record = vulnerability.copy()
                    record.setdefault("package", package.get("name", ""))
                    record.setdefault(
                        "installed_version", package.get("version", "")
                    )
                    vulnerabilities.append(record)

    return [item for item in vulnerabilities if isinstance(item, dict)]


def normalize_vulnerability(vulnerability):
    """Convert a vulnerability record into a clean table row."""
    return {
        "CVE ID": vulnerability.get("cve_id", vulnerability.get("id", "")),
        "Severity": str(vulnerability.get("severity", "UNKNOWN")).upper(),
        "Package": vulnerability.get(
            "package", vulnerability.get("artifact", "")
        ),
        "Installed Version": vulnerability.get(
            "installed_version", vulnerability.get("version", "")
        ),
        "Affected Range": vulnerability.get(
            "affected_range", vulnerability.get("affected_version", "")
        ),
        "Fixed Version": vulnerability.get("fixed_version", ""),
        "Fix Available": vulnerability.get(
            "is_fixable", vulnerability.get("fix_available", False)
        ),
        "CVSS Score": vulnerability.get(
            "cvss_score", vulnerability.get("security-severity", "")
        ),
        "EPSS Score": vulnerability.get("epss_score", ""),
        "EPSS Percentile": vulnerability.get("epss_percentile", ""),
        "Advisory URL": vulnerability.get("advisory_url", ""),
        "Location": vulnerability.get("location", ""),
    }


def _format_sheet(worksheet, table_name, widths, severity_column=None):
    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 30

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = body_alignment

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    if severity_column:
        severity_colors = {
            "CRITICAL": "F4CCCC",
            "HIGH": "FCE5CD",
            "MEDIUM": "FFF2CC",
            "LOW": "D9EAD3",
        }
        for cell in worksheet[severity_column][1:]:
            fill_color = severity_colors.get(str(cell.value).upper())
            if fill_color:
                cell.fill = PatternFill("solid", fgColor=fill_color)


def write_recommendations_sheet(workbook, recommendations_file):
    """Write structured Docker Scout recommendations to a worksheet."""
    parsed = load_recommendations(recommendations_file)
    worksheet = workbook.create_sheet("Recommendations")
    worksheet.sheet_view.showGridLines = False

    title_fill = PatternFill("solid", fgColor="17365D")
    header_fill = PatternFill("solid", fgColor="5B9BD5")
    header_font = Font(color="FFFFFF", bold=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.merge_cells("A1:B1")
    worksheet["A1"] = "Docker Scout Base Image Recommendations"
    worksheet["A1"].fill = title_fill
    worksheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    worksheet["A1"].alignment = Alignment(horizontal="center")
    worksheet.row_dimensions[1].height = 28

    overview_header_row = 2
    worksheet.cell(overview_header_row, 1, "Field")
    worksheet.cell(overview_header_row, 2, "Value")
    for cell in worksheet[overview_header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    overview_rows = list(parsed["overview"].items())
    for row_number, (field, value) in enumerate(overview_rows, start=3):
        worksheet.cell(row_number, 1, field)
        worksheet.cell(row_number, 2, value)

    overview_end_row = overview_header_row + max(len(overview_rows), 1)
    for row in worksheet.iter_rows(
        min_row=overview_header_row + 1,
        max_row=overview_end_row,
        min_col=1,
        max_col=2,
    ):
        for cell in row:
            cell.alignment = body_alignment

    candidate_title_row = overview_end_row + 3
    worksheet.merge_cells(start_row=candidate_title_row, start_column=1, end_row=candidate_title_row, end_column=9)
    worksheet.cell(candidate_title_row, 1, "Candidate Base Images")
    worksheet.cell(candidate_title_row, 1).fill = title_fill
    worksheet.cell(candidate_title_row, 1).font = Font(color="FFFFFF", bold=True)

    candidate_header_row = candidate_title_row + 1
    candidate_headers = [
        "Recommendation",
        "Tag",
        "Details",
        "Pushed",
        "Critical",
        "High",
        "Medium",
        "Low",
        "Unspecified",
    ]
    for column_number, header in enumerate(candidate_headers, start=1):
        cell = worksheet.cell(candidate_header_row, column_number, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    candidate_rows = parsed["candidates"] or [
        {
            "Recommendation": "Status",
            "Tag": "",
            "Details": "No structured candidate images were returned.",
            "Pushed": "",
            "Critical": "",
            "High": "",
            "Medium": "",
            "Low": "",
            "Unspecified": "",
        }
    ]
    for row_number, candidate in enumerate(
        candidate_rows, start=candidate_header_row + 1
    ):
        for column_number, header in enumerate(candidate_headers, start=1):
            cell = worksheet.cell(
                row_number, column_number, candidate.get(header, "")
            )
            cell.alignment = body_alignment

    candidate_end_row = candidate_header_row + len(candidate_rows)
    widths = [24, 28, 58, 16, 12, 10, 10, 10, 14]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    worksheet.freeze_panes = f"A{candidate_header_row + 1}"
    worksheet.auto_filter.ref = (
        f"A{candidate_header_row}:I{candidate_end_row}"
    )


def create_excel_report(
    vulnerabilities, output_file, image_ref, recommendations_file
):
    rows = [normalize_vulnerability(item) for item in vulnerabilities]
    df = pd.DataFrame(rows, columns=REPORT_COLUMNS)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Vulnerabilities", index=False)
        vulnerability_sheet = writer.sheets["Vulnerabilities"]
        _format_sheet(
            vulnerability_sheet,
            "VulnerabilitiesTable",
            [18, 14, 26, 20, 22, 18, 15, 12, 12, 16, 48, 42],
            severity_column=2,
        )

        if not df.empty:
            severity_counts = (
                df["Severity"]
                .value_counts()
                .rename_axis("Severity")
                .reset_index(name="Count")
            )
            severity_order = {
                "CRITICAL": 0,
                "HIGH": 1,
                "MEDIUM": 2,
                "LOW": 3,
                "UNSPECIFIED": 4,
                "UNKNOWN": 5,
            }
            severity_counts["sort"] = severity_counts["Severity"].map(
                lambda value: severity_order.get(value, 99)
            )
            severity_counts = severity_counts.sort_values("sort").drop(
                columns="sort"
            )
        else:
            severity_counts = pd.DataFrame(columns=["Severity", "Count"])

        severity_counts.to_excel(writer, sheet_name="Summary", index=False)
        _format_sheet(writer.sheets["Summary"], "SummaryTable", [20, 12])

        metadata = pd.DataFrame(
            {
                "Field": [
                    "Image",
                    "Scan Date",
                    "Total Vulnerabilities",
                    "Critical",
                    "High",
                    "Medium",
                    "Low",
                ],
                "Value": [
                    image_ref,
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    len(df),
                    len(df[df["Severity"] == "CRITICAL"]),
                    len(df[df["Severity"] == "HIGH"]),
                    len(df[df["Severity"] == "MEDIUM"]),
                    len(df[df["Severity"] == "LOW"]),
                ],
            }
        )
        metadata.to_excel(writer, sheet_name="Scan Info", index=False)
        _format_sheet(writer.sheets["Scan Info"], "ScanInfoTable", [28, 42])

        write_recommendations_sheet(writer.book, recommendations_file)

    print(f"Security report generated: {output_file}")


def main():
    parser = argparse.ArgumentParser(
        description="Convert a vulnerability scan report into an Excel workbook."
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=Path(os.getenv("VULNERABILITY_REPORT_INPUT", DEFAULT_INPUT_FILE)),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(os.getenv("VULNERABILITY_REPORT_OUTPUT", DEFAULT_OUTPUT_FILE)),
    )
    parser.add_argument(
        "--recommendations",
        type=Path,
        default=Path(
            os.getenv(
                "DOCKER_SCOUT_RECOMMENDATIONS",
                DEFAULT_RECOMMENDATIONS_FILE,
            )
        ),
    )
    parser.add_argument(
        "--image",
        default=os.getenv("IMAGE_REF", "portfolio-site-secure:local"),
    )
    args = parser.parse_args()

    print("Loading vulnerability report...")
    data = load_report(args.input)
    vulnerabilities = extract_vulnerabilities(data)
    print(f"Found {len(vulnerabilities)} vulnerabilities.")
    create_excel_report(
        vulnerabilities,
        args.output,
        args.image,
        args.recommendations,
    )


if __name__ == "__main__":
    main()
